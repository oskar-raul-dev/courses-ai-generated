"""La presentación del comité mensual: una diapositiva por sede, con tabla y gráfico nativos."""

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

SEDES = {"Centro": (187_450_000, 162_300_000, 3_100_000),
         "Suba": (142_900_000, 118_600_000, 9_800_000),
         "Kennedy": (98_300_000, 91_200_000, 1_200_000)}


def pesos(v: int) -> str:
    return "$" + f"{v:,}".replace(",", ".")


deck = Presentation()                          # en la casa: Presentation("plantilla-aurea.pptx")
title_only = deck.slide_layouts[5]

summary = deck.slides.add_slide(deck.slide_layouts[0])
summary.shapes.title.text = "Comité mensual · septiembre de 2026"
summary.placeholders[1].text = "Red Áurea: facturado, recaudo y mora por sede"

for sede, (billed, collected, overdue) in SEDES.items():
    slide = deck.slides.add_slide(title_only)
    slide.shapes.title.text = f"Sede {sede}"
    rows = [("Facturado", pesos(billed)), ("Recaudado", pesos(collected)),
            ("Recaudo", f"{collected / billed:.1%}".replace(".", ",")), ("Mora > 90 días", pesos(overdue))]
    table = slide.shapes.add_table(len(rows), 2, Inches(0.5), Inches(1.6), Inches(4.5), Inches(2)).table
    for i, (label, value) in enumerate(rows):
        table.cell(i, 0).text, table.cell(i, 1).text = label, value
        table.cell(i, 1).text_frame.paragraphs[0].runs[0].font.size = Pt(16)

    chart_data = CategoryChartData()
    chart_data.categories = ["Facturado", "Recaudado"]
    chart_data.add_series("Pesos", (billed, collected))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5.3), Inches(1.6),
                           Inches(4.2), Inches(3.5), chart_data)

deck.save("comite-2026-09.pptx")

# Verificación: abrir el archivo y leer lo que quedó, como lo leería PowerPoint.
for n, slide in enumerate(Presentation("comite-2026-09.pptx").slides, 1):
    kinds = sorted(s.shape_type.name for s in slide.shapes if not s.is_placeholder)
    print(n, slide.shapes.title.text, kinds)
