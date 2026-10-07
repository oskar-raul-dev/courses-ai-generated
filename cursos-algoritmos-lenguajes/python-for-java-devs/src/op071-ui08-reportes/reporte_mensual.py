"""El reporte mensual de cartera en PDF (WeasyPrint) y en Excel que suma (XlsxWriter)."""

from jinja2 import Environment
from openpyxl import load_workbook
from weasyprint import HTML
import xlsxwriter

MES = "septiembre de 2026"
ROWS = [("Centro", 187_450_000, 162_300_000, 3_100_000), ("Suba", 142_900_000, 118_600_000, 9_800_000),
        ("Kennedy", 98_300_000, 91_200_000, 1_200_000), ("Zipaquirá", 61_750_000, 52_400_000, 4_450_000)]

# ------------------------------------------------- PDF: plantilla + CSS de impresión
env = Environment(autoescape=True)
env.filters["pesos"] = lambda v: "$" + f"{v:,}".replace(",", ".")
PAGE = env.from_string("""<!doctype html><html lang="es"><meta charset="utf-8"><style>
@page { size: Letter; margin: 2cm; @bottom-right { content: "Página " counter(page) " de " counter(pages); } }
body { font-family: sans-serif; font-size: 10pt; }
td.n { text-align: right; } .alerta { color: #b00020; font-weight: bold; }
</style><h1>Cartera de la red · {{ mes }}</h1><table>
<tr><th>Sede</th><th>Facturado</th><th>Recaudado</th><th>Mora &gt; 90 días</th></tr>
{% for sede, fac, rec, mora in rows %}<tr><td>{{ sede }}</td><td class="n">{{ fac | pesos }}</td>
<td class="n">{{ rec | pesos }}</td><td class="n{{ ' alerta' if mora > 4000000 }}">{{ mora | pesos }}</td></tr>
{% endfor %}</table></html>""")
document = HTML(string=PAGE.render(mes=MES, rows=ROWS)).render()
document.write_pdf("cartera-2026-09.pdf")
print("PDF:", len(document.pages), "página(s)")

# ------------------------------------------------- Excel: números con formato, y una fórmula
book = xlsxwriter.Workbook("cartera-2026-09.xlsx")
sheet = book.add_worksheet("Cartera")
money = book.add_format({"num_format": '"$"#,##0'})
percent = book.add_format({"num_format": "0.0%"})
bold = book.add_format({"bold": True})
sheet.write_row(0, 0, ["Sede", "Facturado", "Recaudado", "Mora > 90 días", "Recaudo"], bold)
for r, (sede, fac, rec, mora) in enumerate(ROWS, start=1):
    sheet.write(r, 0, sede)
    sheet.write_number(r, 1, fac, money)
    sheet.write_number(r, 2, rec, money)
    sheet.write_number(r, 3, mora, money)
    sheet.write_formula(r, 4, f"=C{r + 1}/B{r + 1}", percent)
total = len(ROWS) + 1
sheet.write(total, 0, "Total", bold)
sheet.write_formula(total, 1, f"=SUM(B2:B{total})", money)
sheet.conditional_format(1, 3, len(ROWS), 3, {"type": "cell", "criteria": ">", "value": 4_000_000,
                                              "format": book.add_format({"font_color": "#B00020"})})
sheet.set_column(0, 0, 12)
sheet.set_column(1, 4, 16)
sheet.autofilter(0, 0, len(ROWS), 4)
sheet.freeze_panes(1, 0)
book.close()

# Verificación: lo que verá Excel en la celda de Suba.
cell = load_workbook("cartera-2026-09.xlsx")["Cartera"]["B3"]
print("Excel B3:", cell.value, type(cell.value).__name__, cell.number_format)
