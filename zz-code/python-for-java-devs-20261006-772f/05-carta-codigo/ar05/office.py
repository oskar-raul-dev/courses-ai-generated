"""Office desde Python: Excel con dos bibliotecas, la fórmula que nadie calculó, Word con plantilla y una diapositiva."""

import os
import time
import tracemalloc

import openpyxl
import xlsxwriter
from docx import Document
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

BRANCHES = ["Centro", "Chapinero", "Suba", "Kennedy", "Usaquén", "Engativá", "Fontibón", "Restrepo"]
N = 100_000


def measure(fn):
    start = time.perf_counter(); fn(); elapsed = time.perf_counter() - start
    tracemalloc.start(); fn(); peak = tracemalloc.get_traced_memory()[1]; tracemalloc.stop()   # aparte: tracemalloc frena
    return f"{elapsed:5.2f} s · memoria máx. {peak / 2**20:5.0f} MB"


# 1) 100.000 filas con openpyxl y con XlsxWriter (modo de memoria constante)
def with_openpyxl():
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["fila", "sede", "valor"])
    for i in range(N):
        ws.append([i, BRANCHES[i % 8], i * 1.5])
    ws.append(["", "total", f"=SUM(C2:C{N + 1})"])
    wb.save("openpyxl.xlsx")


def with_xlsxwriter():
    wb = xlsxwriter.Workbook("xlsxwriter.xlsx", {"constant_memory": True}); ws = wb.add_worksheet()
    ws.write_row(0, 0, ["fila", "sede", "valor"])
    for i in range(N):
        ws.write_row(i + 1, 0, [i, BRANCHES[i % 8], i * 1.5])
    ws.write_row(N + 1, 1, ["total", f"=SUM(C2:C{N + 1})"])
    wb.close()


print(f"openpyxl   {measure(with_openpyxl)} · {os.path.getsize('openpyxl.xlsx') / 2**20:.1f} MB")
print(f"XlsxWriter {measure(with_xlsxwriter)} · {os.path.getsize('xlsxwriter.xlsx') / 2**20:.1f} MB")

# 2) La fórmula: openpyxl no calcula
total = openpyxl.load_workbook("openpyxl.xlsx").active.cell(N + 2, 3).value
cached = openpyxl.load_workbook("openpyxl.xlsx", data_only=True).active.cell(N + 2, 3).value
print(f"la celda del total: {total!r} · con data_only=True: {cached!r}")

# 3) Word: una plantilla con el marcador partido en dos "runs" por un cambio de formato
template = Document()
paragraph = template.add_paragraph("Informe de la sede ")
paragraph.add_run("{{se")
paragraph.add_run("de}}").bold = True                     # alguien puso en negrita la mitad del marcador
template.save("plantilla.docx")
doc = Document("plantilla.docx")
for p in doc.paragraphs:
    for run in p.runs:
        run.text = run.text.replace("{{sede}}", "Kennedy")
print(f"reemplazo run por run: {doc.paragraphs[0].text!r} · runs: {[r.text for r in doc.paragraphs[0].runs]}")

# 4) PowerPoint: una diapositiva con un gráfico nativo (editable en PowerPoint)
deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[5])
slide.shapes.title.text = "Filas por sede (muestra)"
data = CategoryChartData(); data.categories = BRANCHES; data.add_series("filas", [N // 8] * 8)
slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(9), Inches(5), data)
deck.save("muestra.pptx")
print(f"muestra.pptx {os.path.getsize('muestra.pptx') / 1024:.0f} KB con un gráfico nativo")
for f in ("openpyxl.xlsx", "xlsxwriter.xlsx", "plantilla.docx", "muestra.pptx"):
    os.remove(f)
